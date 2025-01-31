from flask import Flask, render_template, request, redirect, url_for, session, jsonify

from datetime import datetime, timedelta
import pandas as pd
import os
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment


import platform
import subprocess

import os.path
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

from rapidfuzz import process, fuzz
import unicodedata



SCOPES = ["https://www.googleapis.com/auth/calendar"]

# Dictionnaire des chauffeurs avec leurs identifiants de calendrier
chauffeurs_calendars = {
    "Chauffeur 1": "b304705d464df259f4ac5386ca60a84e80369f4c5addcec12cc9a17ee6b69017@group.calendar.google.com",
    "Chauffeur 2": "dacdd5c9691cddf8feeef9b4897752b788f838ed4918982f65df8c250c61cace@group.calendar.google.com",
    "Chauffeur 3": "1d37cc0687ce0462b489b0ef31dd4cc0a9dd01cfcda5dd0a4f816b5705265b2b@group.calendar.google.com",
    "Chauffeur 4": "1845cd7fba2e73d44394bab7a0d9c7c036fff9e8ee855783d44625aae74a65e5@group.calendar.google.com",
    "Chauffeur 5": "a73d79a42fdcdfa87c248b94d1052c78955ece9a5d59a3f25709ef8ad82a3bd3@group.calendar.google.com",
}
app = Flask(__name__)
app.secret_key = "votre_cle_secrete_pour_la_session"  # Clé secrète pour gérer les sessions

def get_credentials():
    creds = None
    if os.path.exists("token.json"):
        try:
            creds = Credentials.from_authorized_user_file("token.json", SCOPES)
            print("Credentials loaded from file.")
        except Exception as e:
            print(f"Error loading credentials: {e}")

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            try:
                creds.refresh(Request())
                print("Credentials refreshed successfully.")
            except Exception as e:
                print(f"Error refreshing credentials: {e}")
        else:
            print("Authentication required: follow the generated link.")
            flow = InstalledAppFlow.from_client_secrets_file("credentials.json", SCOPES)
            creds = flow.run_local_server(port=0)

        with open("token.json", "w") as token:
            token.write(creds.to_json())
            print("New credentials saved to token.json.")

    return creds


@app.route('/refresh-credentials')
def refresh_credentials():
    get_credentials()
    return "Credentials refreshed successfully!"



# Charger la base de données clients
def load_data(file_path):
    try:
        df = pd.read_excel(file_path, engine='openpyxl')
        df['Nom'] = df['Nom'].str.strip().str.upper()  # Supprime les espaces et convertit en majuscules
        df['Prénom'] = df['Prénom'].str.strip().str.upper()  #
        df['Montant_trimestre_1'].fillna(0, inplace=True)
        df['Nb_courses_trimestre_1'].fillna(0, inplace=True)
        df['Montant_trimestre_2'].fillna(0, inplace=True)
        df['Nb_courses_trimestre_2'].fillna(0, inplace=True)
        df['Montant_trimestre_3'].fillna(0, inplace=True)
        df['Nb_courses_trimestre_3'].fillna(0, inplace=True)
        return df
    except Exception as e:
        print(f"Erreur lors du chargement du fichier : {e}")
        return None

# Fonction pour normaliser un texte (supprimer les accents et mettre en majuscules)
def normalize(text):
    if isinstance(text, str):
        return "".join(c for c in unicodedata.normalize('NFKD', text) if not unicodedata.combining(c)).upper().strip()
    return ""

def get_client_info(nom, prenom, file_path="Base_de_donnees.xlsx"):
    df = load_data(file_path)
    if df is not None:
        # Normaliser les noms et prénoms pour éviter les problèmes de majuscules et d'accents
        df['Nom_Normalisé'] = df['Nom'].apply(normalize)
        df['Prénom_Normalisé'] = df['Prénom'].apply(normalize)

        # Normaliser l'entrée utilisateur
        input_name = f"{normalize(nom)} {normalize(prenom)}"
        clients_list = df.apply(lambda row: f"{row['Nom_Normalisé']} {row['Prénom_Normalisé']}", axis=1).tolist()

        # Rechercher la meilleure correspondance avec Fuzzy Matching
        best_match, score, index = process.extractOne(input_name, clients_list, scorer=fuzz.ratio)

        print(f"Nom saisi: {input_name}")
        print(f"Liste des clients en base: {clients_list}")
        print(f"Meilleure correspondance trouvée: {best_match} avec un score de {score}")

        if score >= 80:  # Seuil de confiance ajustable (80 = bonne tolérance aux erreurs)
            return df.iloc[index]  # Retourne les infos du client correspondant
        else:
            print("Aucun client correspondant trouvé avec un score suffisant.")
            return None
    print("Problème lors du chargement du fichier.")
    return None


# Mettre à jour les informations d'un client après une réservation
def update_client_data(nom, prenom, trimestre, montant_ajoute, file_path="Base_de_donnees.xlsx"):
    df = load_data(file_path)
    if df is not None:
        client = df[(df['Nom'] == nom) & (df['Prénom'] == prenom)]
        if not client.empty:
            client_index = client.index[0]
            if trimestre == 1:
                df.at[client_index, 'Nb_courses_trimestre_1'] += 1
                df.at[client_index, 'Montant_trimestre_1'] += montant_ajoute
            elif trimestre == 2:
                df.at[client_index, 'Nb_courses_trimestre_2'] += 1
                df.at[client_index, 'Montant_trimestre_2'] += montant_ajoute
            elif trimestre == 3:
                df.at[client_index, 'Nb_courses_trimestre_3'] += 1
                df.at[client_index, 'Montant_trimestre_3'] += montant_ajoute
            elif trimestre == 4:
                df.at[client_index, 'Nb_courses_trimestre_4'] += 1
                df.at[client_index, 'Montant_trimestre_4'] += montant_ajoute

            df.to_excel(file_path, index=False, engine='openpyxl')
            print("Données mises à jour avec succès.")
        else:
            print("Client non trouvé.")
    else:
        print("Échec du chargement du fichier.")


def get_montant_trajet(type_trajet):
    prix_trajet = {
        "Asnières Aller": 1.0,
        "Asnières Retour": 1.0,
        "Hors Asnières Aller": 2.0,
        "Hors Asnières Retour": 2.0
    }
    return prix_trajet.get(type_trajet, 0)

# Fonction pour obtenir le trimestre en fonction de la date
def get_trimestre(date):
    if isinstance(date, str):
        date = datetime.strptime(date, '%Y-%m-%d')
    mois = date.month
    return (mois - 1) // 3 + 1

# Fonction pour déduire le montant lors d'une annulation
def deduct_client_data(nom, prenom, trimestre, montant_a_deduire, file_path="Base_de_donnees.xlsx"):
    df = load_data(file_path)
    if df is not None:
        client = df[(df['Nom'] == nom) & (df['Prénom'] == prenom)]
        if not client.empty:
            client_index = client.index[0]
            if trimestre == 1:
                df.at[client_index, 'Nb_courses_trimestre_1'] = max(0, df.at[client_index, 'Nb_courses_trimestre_1'] - 1)
                df.at[client_index, 'Montant_trimestre_1'] = max(0, df.at[client_index, 'Montant_trimestre_1'] - montant_a_deduire)
            elif trimestre == 2:
                df.at[client_index, 'Nb_courses_trimestre_2'] = max(0, df.at[client_index, 'Nb_courses_trimestre_2'] - 1)
                df.at[client_index, 'Montant_trimestre_2'] = max(0, df.at[client_index, 'Montant_trimestre_2'] - montant_a_deduire)
            elif trimestre == 3:
                df.at[client_index, 'Nb_courses_trimestre_3'] = max(0, df.at[client_index, 'Nb_courses_trimestre_3'] - 1)
                df.at[client_index, 'Montant_trimestre_3'] = max(0, df.at[client_index, 'Montant_trimestre_3'] - montant_a_deduire)
            elif trimestre == 4:
                df.at[client_index, 'Nb_courses_trimestre_4'] = max(0, df.at[client_index, 'Nb_courses_trimestre_4'] - 1)
                df.at[client_index, 'Montant_trimestre_4'] = max(0, df.at[client_index, 'Montant_trimestre_4'] - montant_a_deduire)

            df.to_excel(file_path, index=False, engine='openpyxl')
            print(f"Déduction de {montant_a_deduire}€ effectuée avec succès pour le client {nom} {prenom} au trimestre {trimestre}")
            return True
        else:
            print("Client non trouvé.")
            return False
    else:
        print("Échec du chargement du fichier.")
        return False
    


def add_event_to_calendar(chauffeur, summary, description, start_time, end_time):
    """
    Ajoute un événement au calendrier Google du chauffeur spécifié.
    Retourne l'ID de l'événement créé.
    """
    creds = get_credentials()

    try:
        service = build("calendar", "v3", credentials=creds)

        if chauffeur not in chauffeurs_calendars:
            print(f"❌ Erreur : Aucun calendrier configuré pour {chauffeur}")
            return False

        calendar_id = chauffeurs_calendars[chauffeur]

        event = {
            "summary": summary,
            "description": description,
            "start": {
                "dateTime": start_time.isoformat(),
                "timeZone": "Europe/Paris",
            },
            "end": {
                "dateTime": end_time.isoformat(),
                "timeZone": "Europe/Paris",
            },
        }

        event = service.events().insert(calendarId=calendar_id, body=event).execute()
        print(f"✅ Événement ajouté avec succès : {event.get('htmlLink')}")
        return event.get('id')

    except HttpError as error:
        print(f"❌ Une erreur s'est produite : {error}")
        return False


def delete_event_from_calendar(chauffeur, event_id):
    """
    Supprime un événement du calendrier Google du chauffeur spécifié.
    Retourne True si la suppression est réussie, False sinon.
    """
    creds = get_credentials()

    try:
        service = build("calendar", "v3", credentials=creds)

        if chauffeur not in chauffeurs_calendars:
            print(f"❌ Erreur : Aucun calendrier configuré pour {chauffeur}")
            return False

        calendar_id = chauffeurs_calendars[chauffeur]

        service.events().delete(calendarId=calendar_id, eventId=event_id).execute()
        print(f"✅ Événement {event_id} supprimé avec succès du calendrier {chauffeur}.")
        return True

    except HttpError as error:
        print(f"❌ Une erreur s'est produite lors de la suppression de l'événement : {error}")
        return False



def update_course_data(nom, prenom, adresse_depart, adresse_arrivee, horaire, type_trajet, montant, chauffeur, event_id, file_path="Récapitulatif_Courses.xlsx"):
    try:
        if os.path.exists(file_path):
            wb = load_workbook(file_path)
            ws = wb.active
            if ws.max_row == 0 or ws["A1"].value != "ID Google Calendar":
                headers = [
                    "ID Google Calendar",  # Colonne cachée
                    "Identifiant",
                    "Date", 
                    "Horaire", 
                    "Nom", 
                    "Prénom", 
                    "Adresse Départ", 
                    "Adresse Arrivée", 
                    "Type Trajet", 
                    "Montant", 
                    "Chauffeur"
                ]
                ws.delete_rows(1, ws.max_row)
                ws.append(headers)
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Récapitulatif Courses"
            headers = [
                "ID Google Calendar",  # Colonne cachée
                "Identifiant",
                "Date", 
                "Horaire", 
                "Nom", 
                "Prénom", 
                "Adresse Départ", 
                "Adresse Arrivée", 
                "Type Trajet", 
                "Montant", 
                "Chauffeur"
            ]
            ws.append(headers)

        # Masquer la première colonne (A)
        ws.column_dimensions['A'].hidden = True

        # Styles pour l'en-tête
        header_font = Font(bold=True, color="000000")
        header_fill = PatternFill(start_color="FF5757", end_color="FF5757", fill_type="solid")
        for col in ws[1]:
            col.font = header_font
            col.alignment = Alignment(horizontal="center", vertical="center")
            col.fill = header_fill

        # Déterminer le dernier identifiant et l'incrémenter
        last_id = -1  # Commence à -1 pour que la première valeur soit 0000
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
            if row[0].value and str(row[0].value).isdigit():
                last_id = max(last_id, int(row[0].value))  # Trouver le plus grand ID existant
        
        visible_id = str(last_id + 1).zfill(4)  # Formater en 4 chiffres avec des zéros devant


        prenom_formate = prenom.capitalize()
        new_row = [
            event_id,        # ID Google Calendar (caché)
            visible_id,      # Identifiant visible (4 chiffres incrémentés)
            horaire.strftime('%Y-%m-%d'),
            horaire.strftime('%H:%M'),
            nom,
            prenom_formate,
            adresse_depart,
            adresse_arrivee,
            type_trajet,
            f"{montant} €",
            chauffeur,
        ]
        ws.append(new_row)

        # Styles pour les données (alternance de couleurs, alignement)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if row[0].row % 2 == 0:
                fill = PatternFill(start_color="E8F4FF", end_color="E8F4FF", fill_type="solid")
            else:
                fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            for cell in row:
                cell.fill = fill

        # Ajuster la largeur des colonnes visibles
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            if column != 'A':  # Ne pas ajuster la colonne cachée
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column].width = adjusted_width

        wb.save(file_path)
        print("Fichier des courses mis à jour avec succès.")
    except Exception as e:
        print(f"Erreur lors de la mise à jour des courses : {e}")


@app.route('/')
def index():
    if not session.get('logged_in'):
        return redirect(url_for('login'))  # Rediriger vers la page de connexion si non connecté
    return render_template('Page_accueil.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']

        # Exemple de vérification des identifiants (à remplacer par une vraie authentification)
        if username == "admin" and password == "password123":
            session['logged_in'] = True
            return redirect(url_for('index'))  # Rediriger vers l'accueil après connexion
        else:
            return render_template('login.html', error="Nom d'utilisateur ou mot de passe incorrect.")

    return render_template('login.html')

@app.route('/logout')
def logout():
    session.pop('logged_in', None)  # Supprime la session de l'utilisateur
    return redirect(url_for('login'))  # Redirige vers la page de connexion


@app.route('/formulaire')
def formulaire():
    return render_template('formulaire.html')



@app.route('/validation', methods=['POST'])
def validation():
    nom = request.form['nom'].strip().upper()
    prenom = request.form['prenom'].strip().upper()
    adresse_depart = request.form['adresse_depart']
    adresse_arrivee = request.form['adresse_arrivee']
    horaire = request.form['horaire']
    type_trajet = request.form['type_trajet']
    duree_trajet = int(request.form['duree_trajet'])
    chauffeur = request.form['chauffeur']

    if adresse_depart == "autre":
        adresse_depart = request.form['adresse_depart_autre'].strip()
    elif adresse_depart == "domicile":
        client_info = get_client_info(nom, prenom)
        if client_info is not None:
            adresse_depart = client_info['Adresse']
        else:
            return render_template('client_non_trouve.html', nom=nom, prenom=prenom), 404  # On affiche ce qui a été saisi


    if adresse_arrivee == "autre":
        adresse_arrivee = request.form['adresse_arrivee_autre'].strip()
    elif adresse_arrivee == "domicile":
        client_info = get_client_info(nom, prenom)
        if client_info is not None:
            adresse_arrivee = client_info['Adresse']
        else:
            return render_template('client_non_trouve.html', nom=nom, prenom=prenom), 404  # On affiche ce qui a été saisi


    try:
        date_reservation = datetime.strptime(horaire, "%Y-%m-%dT%H:%M")
    except ValueError:
        return "Le format de la date est incorrect.", 400

    heure_arrivee = date_reservation + timedelta(minutes=duree_trajet)
    mois = date_reservation.month
    trimestre = (mois - 1) // 3 + 1

    prix_trajet = {
       "Asnières Aller": 1.0,
       "Asnières Retour": 1.0,
       "Hors Asnières Aller": 2.0,
       "Hors Asnières Retour": 2.0
    }
    montant = prix_trajet.get(type_trajet, 0)

    client_info = get_client_info(nom, prenom)
    if client_info is not None:
        # Ajouter l'événement au calendrier Google
        event_summary = f"Course pour {nom} {prenom}"
        event_description = f"Départ : {adresse_depart}\nArrivée : {adresse_arrivee}\nType : {type_trajet}\nDurée estimée : {duree_trajet} min"
        
        event_id = add_event_to_calendar(chauffeur, event_summary, event_description, date_reservation, heure_arrivee)
        
        if event_id:
            update_client_data(nom, prenom, trimestre, montant)
            update_course_data(nom, prenom, adresse_depart, adresse_arrivee, date_reservation, type_trajet, montant, chauffeur, event_id)
            print("Événement ajouté au calendrier du chauffeur.")
            
            return render_template(
                'validation.html',
                nom=nom,
                prenom=prenom,
                adresse_depart=adresse_depart,
                adresse_arrivee=adresse_arrivee,
                horaire=date_reservation,
                type_trajet=type_trajet,
                montant=montant,
                trimestre=trimestre,
                duree_trajet=duree_trajet,
                chauffeur=chauffeur
            )
        else:
            return "Erreur lors de l'ajout de l'événement au calendrier.", 500
    else:
        return render_template('client_non_trouve.html', nom=nom, prenom=prenom), 404  # On affiche ce qui a été saisi





@app.route('/summary')
def summary():
    df = pd.read_excel('Base_de_donnees.xlsx')
    
    # Get unique names and create client data
    noms = df['Nom'].unique().tolist()
    client_data = df[['Nom', 'Prénom', 
                      'Montant_trimestre_1', 'Montant_trimestre_2', 'Montant_trimestre_3','Montant_trimestre_4',
                      'Nb_courses_trimestre_1', 'Nb_courses_trimestre_2', 'Nb_courses_trimestre_3','Nb_courses_trimestre_4']].to_dict('records')
    
    return render_template('summary.html', noms=noms, client_data=client_data)
    
@app.route('/open_excel', methods=['GET'])
def open_excel():
    
    chemin_fichier = os.path.join(os.getcwd(), 'Récapitulatif_Courses.xlsx')
    
    try:
        # Ouvrir le fichier Excel
        if platform.system() == "Windows":
            os.startfile(chemin_fichier)
        elif platform.system() == "Darwin":  # macOS
            subprocess.run(["open", chemin_fichier])
        elif platform.system() == "Linux":
            subprocess.run(["xdg-open", chemin_fichier])

        return jsonify({"message": "Fichier Excel ouvert avec succès"}), 200
    except Exception as e:
        return jsonify({"error": f"Erreur lors de l'ouverture du fichier : {str(e)}"}), 500




@app.route('/annulation', methods=['GET', 'POST'])
def annulation():
    if request.method == 'GET':
        return render_template('annulation.html')

    try:
        prenom = request.form.get('Prénom', '').strip()
        nom = request.form.get('Nom', '').strip()
        mois = request.form.get('Mois', '').strip()
        file_path = "Récapitulatif_Courses.xlsx"

        if not all([prenom, nom, mois]):
            missing_fields = [field for field in ["Prénom", "Nom", "Mois"] if not request.form.get(field, '').strip()]
            return jsonify({"error": "Champs manquants", "missing": missing_fields}), 400

        if not os.path.exists(file_path):
            return jsonify({"error": "Fichier des courses introuvable"}), 404

        try:
            wb = load_workbook(filename=file_path, read_only=False, data_only=True)
            ws = wb.active
        except Exception as e:
            return jsonify({"error": "Erreur lors de la lecture du fichier Excel", "details": str(e)}), 500

        try:
            mois_obj = datetime.strptime(mois, '%Y-%m').date()
        except ValueError:
            return jsonify({"error": "Format de mois invalide", "details": "Le format attendu est YYYY-MM"}), 400

        reservations = []

        for row in range(2, ws.max_row + 1):
            try:
                row_date = ws.cell(row=row, column=3).value  # Colonne 3 = Date
                if isinstance(row_date, datetime):
                    row_date = row_date.date()
                elif isinstance(row_date, str):
                    try:
                        row_date = datetime.strptime(row_date, '%Y-%m-%d').date()
                    except ValueError:
                        continue

                row_nom = str(ws.cell(row=row, column=5).value or '').strip().upper()
                row_prenom = str(ws.cell(row=row, column=6).value or '').strip().upper()

                if row_date.year == mois_obj.year and row_date.month == mois_obj.month and row_nom == nom.upper() and row_prenom == prenom.upper():
                    reservations.append({
                        "id": ws.cell(row=row, column=1).value,  # ID Google Calendar
                        "date": row_date.strftime('%Y-%m-%d'),
                        "horaire": ws.cell(row=row, column=4).value,
                        "depart": ws.cell(row=row, column=7).value,
                        "arrivee": ws.cell(row=row, column=8).value,
                        "montant": get_montant_trajet(ws.cell(row=row, column=9).value)  # Colonne 9 = Type Trajet
                    })
            except Exception as e:
                print(f"Erreur lors du traitement de la ligne {row}: {str(e)}")
                continue

        if not reservations:
            return jsonify({"message": "Aucune réservation trouvée pour ce mois"}), 200

        return jsonify({
            "success": True,
            "reservations": reservations,
            
        }), 200

    except Exception as e:
        return jsonify({"error": "Erreur inattendue", "details": str(e)}), 500




@app.route('/annuler_course/<event_id>', methods=['POST'])
def annuler_course(event_id):
    try:
        file_path = "Récapitulatif_Courses.xlsx"

        if not os.path.exists(file_path):
            return jsonify({"error": "Fichier des courses introuvable"}), 404

        try:
            wb = load_workbook(filename=file_path, read_only=False, data_only=True)
            ws = wb.active
        except Exception as e:
            return jsonify({"error": "Erreur lors de la lecture du fichier Excel", "details": str(e)}), 500

        row_to_delete = None

        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == event_id:
                row_to_delete = row
                break

        if row_to_delete is None:
            return jsonify({"error": "Réservation introuvable"}), 404

        try:
            row_date = ws.cell(row=row_to_delete, column=3).value
            if isinstance(row_date, datetime):
                row_date = row_date.date()
            elif isinstance(row_date, str):
                row_date = datetime.strptime(row_date, '%Y-%m-%d').date()

            nom = ws.cell(row=row_to_delete, column=5).value.strip().upper()
            prenom = ws.cell(row=row_to_delete, column=6).value.strip().upper()
            type_trajet = ws.cell(row=row_to_delete, column=9).value
            chauffeur = ws.cell(row=row_to_delete, column=11).value

            montant = get_montant_trajet(type_trajet)
            trimestre = get_trimestre(row_date)

            deduction_success = deduct_client_data(nom, prenom, trimestre, montant)
            if not deduction_success:
                return jsonify({"error": "Échec de la déduction du montant"}), 500

            if chauffeur and event_id:
                event_deleted = delete_event_from_calendar(chauffeur, event_id)
                if not event_deleted:
                    return jsonify({"error": f"Échec de la suppression de l'événement Google Calendar {event_id}"}), 500

            ws.delete_rows(row_to_delete)

            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=2, value=str(row - 2).zfill(4))

            apply_styles(ws)
            wb.save(file_path)

            return jsonify({"success": True, "message": "Réservation annulée avec succès"}), 200

        except Exception as e:
            return jsonify({"error": "Erreur lors de l'annulation de la réservation", "details": str(e)}), 500

    except Exception as e:
        return jsonify({"error": "Erreur inattendue", "details": str(e)}), 500
    
    
def apply_styles(ws):
    """Applique les styles au worksheet"""
    # Style d'en-tête
    header_style = {
        'font': Font(bold=True, color="000000"),
        'fill': PatternFill(start_color="FF5757", end_color="FF5757", fill_type="solid"),
        'alignment': Alignment(horizontal="center", vertical="center")
    }
    
    for cell in ws[1]:
        cell.font = header_style['font']
        cell.fill = header_style['fill']
        cell.alignment = header_style['alignment']

    # Style des données
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        # Alignement
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Alternance de couleurs
        fill = PatternFill(
            start_color="E8F4FF" if row[0].row % 2 == 0 else "FFFFFF",
            end_color="E8F4FF" if row[0].row % 2 == 0 else "FFFFFF",
            fill_type="solid"
        )
        
        for cell in row:
            cell.fill = fill
            



if __name__ == "__main__":
   
    app.run(debug=True)
